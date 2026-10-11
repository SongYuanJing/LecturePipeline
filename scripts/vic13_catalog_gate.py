"""Disposable native Tk catalog gate. No workers, ASR, network or production config."""
import argparse
import json
import os
from pathlib import Path
import sys
import tempfile
import time
from unittest.mock import patch

REPO=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(REPO/'tests'),str(REPO/'src/app'),str(REPO/'src/app/gui_v1.8')]
from test_lecture_catalog import fixture,add,dump,fingerprints
from pipeline_config import read
from lecture_catalog import build


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--interactive',action='store_true');args=parser.parse_args()
    base=Path(tempfile.mkdtemp(prefix='LP-vic13-gui-'));c=fixture(base)
    for code,key,stage,parts in [('ECON','01_2026-10-01','done',1),('ECON','02_2026-10-03','done',8),('ECON','03_2026-10-05','done',1),
                                  ('MICRO','01_2026-10-02','done',1),('MICRO','02_2026-10-04','done',1),('MICRO','03_2026-10-06','waiting_ai',1),
                                  ('RESEARCH','01_2026-10-07','raw',1),('RESEARCH','02_2026-10-08','done',1)]:add(c,code,key,stage,parts)
    # Real openable DOCX files; processing authorities remain explicitly synthetic.
    from docx import Document
    for p in c.workspace_root.glob('**/lecture.docx'):
        d=Document();d.add_heading('Disposable catalog test — 研究方法',0);d.add_paragraph('Тестовая лекция. Реальных учебных данных нет.');d.save(p)
    (c.subjects()['RESEARCH']['ready_dir']/'02_2026-10-08/lecture.docx').unlink()
    sources={str(n):dict(source_id=str(n),course_id='ECON',source_path=str(base/f'20261010_{n}.wav'),part_hint=n,status='review',group_id='group',lecture_date='2026-10-10',lecture_date_source='filename') for n in (1,3)}
    dump(c.worker_home/'inbox.json',dict(schema_version=1,sources=sources,groups={'group':dict(group_id='group',course_id='ECON',session_stem='20261010',member_ids=['1','3'],status='review',error='Не хватает части 2',lecture_date='2026-10-10',lecture_date_source='filename')},courses={}))
    dump(c.worker_home/'worker_status.json',dict(worker='STOPPED',lectures={}))
    dump(c.dialogue_home/'status.json',dict(worker='STOPPED',files={}))
    dump(c.dialogue_home/'dialogue_state.json',dict(jobs={}))
    (c.dialogue_root/'00 — Входящие аудио').mkdir(parents=True)
    before={**fingerprints(c.data_root),'CONFIG':(c.path.read_bytes().hex(),c.path.stat().st_mtime_ns)}
    service_files={p:(p.read_bytes(),p.stat().st_mtime_ns) for root in (c.worker_home,c.queue_home) for p in root.rglob('*') if p.is_file()}
    dispatches=[];real_startfile=os.startfile
    def observed_startfile(path):
        real_startfile(path)  # Delegate to the real Windows API; no recording-only opener.
        dispatches.append(str(path));dump(base/'windows-open.json',dict(dispatched=dispatches))
    first=build(c);start=time.perf_counter();assert first==build(c);refresh=time.perf_counter()-start
    import tkinter as tk
    import adapter
    from app import App
    from dpi import enable_system_dpi
    enable_system_dpi();root=tk.Tk();os.environ['LECTURE_CONFIG']=str(c.path)
    class Controls:
        def __init__(self,c):pass
        def states(self):return {k:dict(running=False,enabled=False) for k in ('lecture','dialogue')}
    errors=[];root.report_callback_exception=lambda typ,value,tb:errors.append(str(value))
    with patch.object(adapter,'Controls',Controls),patch.object(adapter,'preflight',lambda c:dict(ok=True,checks=[])),patch.object(os,'startfile',observed_startfile):
        app=App(root,c.path);root.title('VIC-13 disposable catalog · 200%');app.tabs.select(app.pages['Лекции'])
        report=dict(base=str(base),fixture=True,refresh_seconds=refresh)
        def check():
            if app.busy:root.after(100,check);return
            try:
                assert app.view is not None and 'catalog' in app.view,app.banner.get()
                panel=app.catalog;assert len(panel.rows)==8
                assert len(panel.pending_rows)==1
                assert round(float(root.tk.call('tk','scaling'))*72)==192
                panel.course.set('Эконометрика [ECON]');panel.fill();assert len(panel.rows)==3
                panel.status.set('Готово');panel.fill();assert len(panel.rows)==3
                panel.course.set('Все курсы');panel.status.set('Все состояния');panel.fill()
                iid=next(i for i,r in panel.rows.items() if r['parts']==8)
                panel.table.selection_set(iid);panel.select();assert '8.' in panel.detail.get('1.0','end')
                panel.pending.selection_set('0');panel.review();assert app.tabs.select()==str(app.pages['Курсы / Inbox'])
                assert app.course_panel.items.selection()==('group',)
                app.tabs.select(app.pages['Лекции']);panel.select()
                normal=(root.winfo_width(),root.winfo_height());root.geometry('1900x1240');root.update()
                panel.table.xview_moveto(1);root.update_idletasks();assert panel.table.xview()[0]>0
                panel.table.xview_moveto(0)
                for widget in (panel.course_box,panel.buttons['word'],panel.pending):
                    assert widget.winfo_ismapped() and widget.winfo_rooty()+widget.winfo_height()<=root.winfo_rooty()+root.winfo_height()
                root.geometry(f'{normal[0]}x{normal[1]}');root.update()
                # Exact existing safe-open helper, including rejected absent paths.
                opened=[];row=panel.rows[iid]
                for key in ('word','raw','result','source'):adapter.safe_open(c,row['paths'][key],opener=opened.append)
                assert len(opened)==4
                try:adapter.safe_open(c,base/'does-not-exist')
                except (ValueError,FileNotFoundError):pass
                else:raise AssertionError('Missing path accepted')
                outside=base/'outside.txt';outside.write_text('Disposable guard test',encoding='utf8')
                invalid=c.data_root/'invalid.exe';invalid.write_bytes(b'not executable')
                for path in (outside,invalid,c.data_root/'missing.txt'):
                    try:adapter.safe_open(c,path,opener=opened.append)
                    except (ValueError,FileNotFoundError):pass
                    else:raise AssertionError('Unsafe path accepted: '+str(path))
                invalid.unlink();assert len(opened)==4
                after={**fingerprints(c.data_root),'CONFIG':(c.path.read_bytes().hex(),c.path.stat().st_mtime_ns)}
                assert before==after and all((p.read_bytes(),p.stat().st_mtime_ns)==value for p,value in service_files.items())
                assert build(c)==first and not errors,errors
                report.update(status='PASS',lectures=8,pending=1,dpi=192,filters=True,selection_details=True,scroll=True,review_link=True,resize=True,safe_open_guard=True,
                              config_data_sha_mtime_unchanged=True,rebuild_identical=True,errors=errors,os_open='manual visual phase pending')
                dump(base/'evidence.json',report);print(json.dumps(report,ensure_ascii=False),flush=True)
            except Exception as exc:
                import traceback
                report.update(status='FAIL',error=traceback.format_exc(),errors=errors);dump(base/'evidence.json',report);print(json.dumps(report,ensure_ascii=False),flush=True)
                app.close();return
            if not args.interactive:app.close()
        root.after(300,check);root.mainloop()
    if report.get('status')=='PASS':
        after={**fingerprints(c.data_root),'CONFIG':(c.path.read_bytes().hex(),c.path.stat().st_mtime_ns)}
        assert before==after,'Data/config changed during interactive inspection'
        assert all((p.read_bytes(),p.stat().st_mtime_ns)==value for p,value in service_files.items()),'Journal changed'
        assert build(c)==first and not errors,errors
        report.update(real_os_startfile_dispatches=dispatches,post_interaction_invariance=True,
                      os_open='real Windows dispatch recorded' if dispatches else 'not exercised')
        dump(base/'evidence.json',report);print(json.dumps(report,ensure_ascii=False),flush=True)
    assert report.get('status')=='PASS',report


if __name__=='__main__':main()
