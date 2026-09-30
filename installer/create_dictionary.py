"""New empty v1.4-compatible workbook; never migrates or opens user workbooks."""
import io,json,sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table,TableStyleInfo
def create(path,app):
 if path.exists():raise FileExistsError(path)
 sys.path.insert(0,str(app));import dictionary_merge as d
 w=Workbook();s=w.active;s.title=d.MAIN;s['A1']='Vocabulary';s.append([]);s.append([])
 for i,v in enumerate(d.HEAD,1):s.cell(4,i,v)
 s.cell(5,1,'');s.freeze_panes='B5';s.column_dimensions['H'].hidden=True
 for col,width in zip('ABCDEFG',(14,24,24,36,24,14,40)):s.column_dimensions[col].width=width
 dv=DataValidation(type='list',formula1='"Не выучено,Выучено"',allow_blank=True);s.add_data_validation(dv);dv.add('A5:A1048576')
 table=Table(displayName='VocabularyBank',ref='A4:H5');table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);s.add_table(table)
 prov=w.create_sheet(d.PROV);prov.append(['entry_id','kind','event_id','chinese','pinyin','translation','count','subject','lecture_id','source_file_id','signature','source_file_name','original']);prov.sheet_state='hidden'
 meta=w.create_sheet(d.META);meta.append(['key','value']);meta.append(['schema_version','1.4']);meta.append(['baseline_ai_state',json.dumps({'schema_version':1,'processed':{}})]);meta.sheet_state='hidden'
 b=io.BytesIO();w.save(b);w.close();d.Book(b.getvalue()).validate();path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(b.getvalue())
if __name__=='__main__':create(Path(sys.argv[1]),Path(sys.argv[2]))
