"""Local extractive Quick mode. No model, network, speaker inference or lecture imports."""
import argparse,json,re,collections
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn

def stamp(t):
 n=int(round(t*1000));return f'{n//3600000:02d}:{n//60000%60:02d}:{n//1000%60:02d}.{n%1000:03d}'
def uncertain(s):
 t=s['text'];return s.get('avg_logprob',0)<-1 or s.get('no_speech_prob',0)>.6 or bool(re.search(r'(.{3,20})\1\1',t)) or bool(re.search('请不吝点赞|字幕志愿者|请使用简体中文|thanks for watching|subscribe to',t,re.I))
def select(rows):
 clean=[];seen=set()
 for s in rows:
  t=s['text'].strip();key=re.sub(r'\W','',t).lower()
  if not uncertain(s) and len(key)>4 and key not in seen:clean.append(s);seen.add(key)
 def tokens(t):return re.findall(r'[a-z]{3,}|[\u3400-\u9fff]{2}',t.lower())
 counts=collections.Counter(w for s in clean for w in tokens(s['text']))
 scores=lambda s:sum(min(counts[w],5) for w in tokens(s['text']))/(max(1,len(tokens(s['text'])))**.5)
 chosen=sorted(sorted(clean,key=scores,reverse=True)[:3],key=lambda s:s['start'])
 conclusions=[s for s in clean if re.search(r'所以|因此|结论|决定|同意|确认|\b(therefore|conclusion|decided|agreed|confirmed)\b',s['text'],re.I)][:5]
 actions=[s for s in clean if re.search(r'请(?:你|您)?(?:在|把|提交|发送|带|准备|完成)|我会|我们(?:约定|决定|将)|\b(?:I will|we will|I’ll|we’ll|please (?:send|bring|submit|prepare|finish|call)|let us meet|let.s meet)\b',s['text'],re.I) and not re.search(r'不会|不要|不需要|\b(?:not|never|won.t|would|could|if)\b',s['text'],re.I)][:12]
 return chosen,conclusions,actions
def build(request,output):
 r=json.loads(Path(request).read_text(encoding='utf8'));rows=r['segments'];d=Document();sec=d.sections[0]
 sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
 for style in d.styles:
  for n in list(style.element.iter(qn('w:pBdr'))):n.getparent().remove(n)
 for name in ('Normal','Title','Subtitle','Heading 1'):
  st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 d.styles['Normal'].font.size=Pt(11);d.styles['Title'].font.size=Pt(21);d.styles['Heading 1'].font.size=Pt(14)
 d.styles['Normal'].paragraph_format.space_after=Pt(5)
 d.styles['Normal'].paragraph_format.line_spacing=1.05
 d.styles['Heading 1'].paragraph_format.space_before=Pt(14)
 d.add_paragraph('Запись разговора','Title');d.add_paragraph(r['title'],'Subtitle')
 d.add_paragraph('Обработано '+r['processed_at']+' · язык '+r['detected_language'])
 d.add_paragraph('Локальный Quick mode. Обзор, выводы и поручения — выбранные реплики из автоматической расшифровки, без генеративного пересказа. Участники разговора автоматически не определяются.')
 def quote(s,context=False):
  if context and r['detected_language']=='en' and not re.search(r'[.!?]$',s['text'].strip()):
   index=rows.index(s)
   if index+1<len(rows):s=dict(start=s['start'],end=rows[index+1]['end'],text=s['text']+' '+rows[index+1]['text'])
  p=d.add_paragraph(f"[{stamp(s['start'])}–{stamp(s['end'])}] {s['text']}");p.paragraph_format.keep_together=True
 chosen,conclusions,actions=select(rows)
 for title,items,empty in [('Краткое содержание разговора',chosen,'Надёжные реплики для краткого обзора не выделены.'),('Ключевые выводы',conclusions,'Явные итоговые формулировки автоматически не обнаружены.')]:
  d.add_paragraph(title,'Heading 1')
  if not items:d.add_paragraph(empty)
  for s in items:quote(s,context=True)
 if actions:
  d.add_paragraph('Договорённости и действия','Heading 1');d.add_paragraph('Прямые формулировки из записи. Ответственные и сроки не добавляются по догадке.')
  for s in actions:quote(s,context=True)
 d.add_paragraph('Полный текст разговора','Heading 1')
 if not rows:d.add_paragraph('Речь не распознана. Проверьте исходную запись; это не доказывает отсутствие речи.')
 for s in rows:quote(s)
 d.add_paragraph('Сомнительные места','Heading 1');doubts=[s for s in rows if uncertain(s)]
 if not doubts:d.add_paragraph('По автоматическим признакам подозрительные фрагменты не выделены. Это не подтверждает безошибочность распознавания.')
 for s in doubts:quote(s)
 sec.footer.paragraphs[0].text='Dialogue Quick 1.6'
 d.save(output)
 return dict(summary_mode='extractive',summary_segments=len(chosen),conclusion_segments=len(conclusions),action_segments=len(actions),uncertain_segments=len(doubts))
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('request');p.add_argument('output');args=p.parse_args();print(json.dumps(build(args.request,args.output)))
