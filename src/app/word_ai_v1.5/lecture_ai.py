"""Lecture Pipeline Word/AI layer 1.5. Explicit one-lecture workflow; no watcher."""
from __future__ import annotations
import argparse, copy, hashlib, json, os, re, shutil, time, uuid
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

import sys
sys.path.insert(0,str(Path(__file__).resolve().parent.parent))
import atomic_publish
import source_provenance as provenance

VERSION='1.5'
ORG=('homework','deadlines','exam','preparation','announcements','dates')
LABELS=('Домашнее задание','Сроки сдачи','Экзамен и контроль','Что подготовить','Объявления','Важные даты')
def sha(data): return hashlib.sha256(data).hexdigest()
def load(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def dump(p,x): Path(p).write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf8')
def utc(): return datetime.now(timezone.utc).isoformat(timespec='seconds').replace('+00:00','Z')
def need(ok,msg):
 if not ok: raise ValueError(msg)
def nonempty(x): return isinstance(x,str) and bool(x.strip())
def local_file(root,name):
 need(nonempty(name) and Path(name).name==name and '/' not in name and '\\' not in name and ':' not in name,'Unsafe input filename')
 p=(root/name).resolve();need(p.parent==root.resolve() and p.is_file(),'Missing input '+name);return p
def millis(s):
 h,m,sec,ms=map(int,re.split('[:,]',s));need(m<60 and sec<60,'Invalid time');return ((h*60+m)*60+sec)*1000+ms
def parse_srt(path,part):
 out=[];previous=-1
 for block in re.split(r'\n\s*\n',path.read_text(encoding='utf-8-sig').strip()):
  lines=block.splitlines();need(len(lines)>=3,'Malformed SRT')
  need(lines[0].isdigit(),'Invalid SRT cue index');n=int(lines[0]);need(n>previous,'Duplicate/unordered cue');previous=n
  match=re.fullmatch(r'(\d{2,}:\d{2}:\d{2},\d{3}) --> (\d{2,}:\d{2}:\d{2},\d{3})',lines[1]);need(match is not None,'Malformed SRT timestamp')
  a,b=match.groups();need(millis(a)<=millis(b),'Reversed SRT time')
  out.append(dict(id=f'{part}:{n}',part=part,cue=n,start=a,end=b,text=' '.join(lines[2:])))
 need(bool(out),'Empty SRT');return out
def read_bundle(manifest):
 manifest=Path(manifest).resolve();m=load(manifest);root=manifest.parent
 for k in ('lecture','subject_code','subject'):need(nonempty(m.get(k)),'Manifest missing '+k)
 need(re.fullmatch(r'\d{2,3}_\d{4}-\d{2}-\d{2}',m['lecture']) is not None,'Invalid lecture key')
 raw=local_file(root,m['combined_txt']);files=[manifest,raw];segments=[]
 need(isinstance(m.get('outputs'),list) and len(m['outputs'])>0,'Manifest outputs missing')
 need(len(m.get('sources',[]))==len(m['outputs']),'Manifest source/output count mismatch')
 for part,o in enumerate(m['outputs'],1):
  need(o['source']==m['sources'][part-1]['name'],'Manifest source order mismatch')
  p=local_file(root,o['srt']);need(p not in files,'Repeated SRT');files.append(p);segments+=parse_srt(p,part)
 rawtext=raw.read_text(encoding='utf-8-sig');body=[]
 for line in rawtext.splitlines():
  if not line.strip() or line.startswith(('===== ЧАСТЬ','Источник:')):continue
  match=re.match(r'^\[[\d:]+ - [\d:]+\]\s*(.*)$',line);need(match is not None,'Unexpected raw line');body.append(match[1])
 compact=lambda s:re.sub(r'\s+','',s)
 need(compact(''.join(body))==compact(''.join(s['text'] for s in segments)),'raw and SRT content differ')
 hashes={str(p):sha(p.read_bytes()) for p in files}
 return dict(manifest=m,manifest_path=str(manifest),raw_path=str(raw),segments=segments,input_hashes=hashes)
def state_read(path):
 data=Path(path).read_bytes();s=json.loads(data.decode('utf-8-sig'))
 need(s.get('schema_version')==1 and isinstance(s.get('processed'),dict),'Unsupported ai_state schema')
 return data,s
def already(s,source,m):
 return source in s['processed'] or any(r.get('source_file_id')==source or (r.get('subject_code')==m['subject_code'] and r.get('lecture_id')==m['lecture']) for r in s['processed'].values())
def prepare(manifest,state,source_id,out,pilot=False,workspace=None):
 b=read_bundle(manifest);data,s=state_read(state);m=b['manifest']
 identity=provenance.local(b,workspace) if source_id is None else dict(source_file_id=source_id)
 identity_key=provenance.key(identity)
 need(pilot or not already(s,identity_key,m),'Already processed lecture; production reprocessing forbidden')
 dest=Path(out);need(not dest.exists(),'Output directory already exists');dest.mkdir(parents=True)
 raw=Path(b['raw_path']);modified=datetime.fromtimestamp(raw.stat().st_mtime,timezone.utc).isoformat(timespec='milliseconds').replace('+00:00','Z')
 src=dict(**identity,source_signature=f'modified={modified};size={raw.stat().st_size}',source_modified_time=modified,source_size_bytes=raw.stat().st_size,source_file_name=raw.name,subject_code=m['subject_code'],subject=m['subject'],lecture_id=m['lecture'])
 b.update(schema_version=1,layer_version=VERSION,source=src,ai_state_path=str(Path(state).resolve()),ai_state_sha256=sha(data),pilot=pilot)
 dump(dest/'source_packet.json',b)
 prompt=Path(__file__).with_name('AI_INSTRUCTIONS_v1.5.md')
 if prompt.exists():shutil.copy2(prompt,dest/prompt.name)
 return b
def resolve_ref(ref,b):
 need(isinstance(ref,dict),'Invalid evidence reference')
 index={s['id']:s for s in b['segments']};a=index.get(ref.get('start'));z=index.get(ref.get('end'))
 need(a is not None and z is not None and a['part']==z['part'] and a['cue']<=z['cue'],'Invalid evidence range')
 selected=[s for s in b['segments'] if s['part']==a['part'] and a['cue']<=s['cue']<=z['cue']]
 if 'quote' in ref:need(nonempty(ref['quote']) and ref['quote'] in ' '.join(s['text'] for s in selected),'Quote is not in SRT')
 return f"Часть {a['part']} · {a['start'].replace(',', '.')}–{z['end'].replace(',', '.')}"
def validate_content(c,b):
 need(c.get('schema_version')==1,'AI content schema must be 1')
 need(c.get('input_hashes')==b['input_hashes'],'AI content does not match input hashes')
 need(nonempty(c.get('title')),'Title required')
 parts=len(b['manifest']['outputs']);reviews=c.get('part_review',[])
 need(sorted(x['part'] for x in reviews)==list(range(1,parts+1)),'Review every recording part exactly once')
 for r in reviews:need(nonempty(r.get('note')) and r.get('status') in ('usable','mixed','unreliable'),'Part review required')
 blocks=[]
 for k in ('summary','notes','definitions','emphasis','uncertainties','formulas'):
  need(isinstance(c.get(k),list),'Missing section '+k)
  if k in ('summary','notes','definitions','emphasis'):need(bool(c[k]),'Empty substantive section '+k)
  blocks+=c[k]
 need(set(c.get('organization',{}))==set(ORG),'All six organizational categories required')
 for k in ORG:
  o=c['organization'][k];need(o.get('status') in ('confirmed','uncertain','not_found'),'Invalid organization status')
  need(isinstance(o.get('items'),list),'Organization items missing')
  need((o['status']=='not_found')== (not o['items']),'Organization status/items inconsistent');blocks+=o['items']
 for block in blocks:
  need(nonempty(block.get('text')),'Empty content block')
  need(block.get('basis') in ('transcript','reconstruction','uncertain'),'Explicit content basis required')
  need(isinstance(block.get('refs'),list) and bool(block['refs']),'Every substantive block needs SRT evidence')
  for ref in block['refs']:resolve_ref(ref,b)
 for x in c['uncertainties']:need(x['basis']=='uncertain','Uncertainty must be labelled')
 for f in c['formulas']:need('math' in f,'Formula must contain structured native math');math_element(f['math'])
 terms=c.get('vocabulary');need(isinstance(terms,list),'Vocabulary list required');seen=set()
 for term in terms:
  for k in ('chinese','pinyin','russian'):need(nonempty(term.get(k)),'Term missing '+k)
  key=(term['chinese'],term.get('sense_id',''));need(key not in seen,'Duplicate vocabulary term');seen.add(key)
  refs=term.get('refs',[]);need(bool(refs),'Vocabulary evidence missing')
  for ref in refs:resolve_ref(ref,b)
  # Counts are deterministic literal transcript occurrences, including ASR repeats; no estimated audio frequency.
  count=sum(s['text'].count(term['chinese']) for s in b['segments'])
  need(count>0,'Term absent from SRT; normalize only after explicit source correction')
 return c
def math_element(x):
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 def e(tag):return OxmlElement('m:'+tag)
 def append_content(n,content):
  node=math_element(content)
  if node.tag==qn('m:e'):
   for ch in list(node):n.append(ch)
  else:n.append(node)
 def group(tag,content):
  n=e(tag);append_content(n,content);return n
 if isinstance(x,str):
  n=e('r');t=e('t');t.set(qn('xml:space'),'preserve');t.text=x;n.append(t);return n
 if isinstance(x,list):
  n=e('e')
  for v in x:append_content(n,v)
  return n
 need(isinstance(x,dict) and len(x)==1,'Invalid math AST')
 op,args=next(iter(x.items()));need(op in ('sub','sup','frac','hat','bar','sum'),'Unknown math operator')
 need(isinstance(args,list) and len(args)==({'sum':3,'hat':1,'bar':1}.get(op,2)),'Invalid math arguments')
 if op in ('sub','sup','frac'):
  n=e({'sub':'sSub','sup':'sSup','frac':'f'}[op]);tags={'sub':('e','sub'),'sup':('e','sup'),'frac':('num','den')}[op]
  for tag,v in zip(tags,args):n.append(group(tag,v))
 elif op in ('hat','bar'):
  n=e('acc');p=e('accPr');ch=e('chr');ch.set(qn('m:val'),'̂' if op=='hat' else '̅');p.append(ch);n.append(p);n.append(group('e',args[0]))
 else:
  n=e('nary');p=e('naryPr');ch=e('chr');ch.set(qn('m:val'),'∑');p.append(ch);n.append(p)
  for tag,v in zip(('sub','sup','e'),args):n.append(group(tag,v))
 return n
def render_word(c,b,dest):
 from docx import Document
 from docx.shared import Inches,Pt,RGBColor
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.8);s.footer_distance=Inches(.3)
 for st in d.styles:
  for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
  st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(6)
 d.styles['Normal'].paragraph_format.line_spacing=1.08
 d.styles['Title'].font.size=Pt(22);d.styles['Heading 1'].font.size=Pt(15);d.styles['Heading 2'].font.size=Pt(12)
 d.add_paragraph(c['title'],'Title');m=b['manifest']
 d.add_paragraph(f"{m['subject']} · {m['lecture']} · {len(m['outputs'])} частей",'Subtitle')
 d.add_paragraph('Учебный конспект для повторения лекции. Таймкоды отсчитываются от начала указанной части. Восстановленные записи формул и сомнительные места отмечены отдельно.')
 def block(x):
  if x.get('heading'):d.add_paragraph(x['heading'],'Heading 2')
  p=d.add_paragraph(x['text'])
  p.paragraph_format.keep_together=True;p.paragraph_format.keep_with_next=True
  if 'math' in x:
   p=d.add_paragraph();om=OxmlElement('m:oMath');node=math_element(x['math'])
   if node.tag==qn('m:e'):
    for ch in list(node):om.append(ch)
   else:om.append(node)
   p._p.append(om)
   p.paragraph_format.keep_with_next=True
  basis={'transcript':'Опора в записи','reconstruction':'Восстановление по контексту','uncertain':'Требует проверки'}[x['basis']]
  p=d.add_paragraph(basis+': '+'; '.join(resolve_ref(r,b) for r in x['refs']))
  p.paragraph_format.space_after=Pt(8)
  for run in p.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(85,85,85)
 d.add_paragraph('Организационные вопросы','Heading 1')
 absent=[]
 for k,label in zip(ORG,LABELS):
  if not c['organization'][k]['items']:absent.append(label.lower());continue
  d.add_paragraph(label,'Heading 2');o=c['organization'][k]
  for x in o['items']:block(x)
 if absent:d.add_paragraph('В доступной расшифровке достоверно не установлены: '+', '.join(absent)+'. Это не означает, что преподаватель ничего не задавал: повреждённые места требуют проверки.')
 for k,label in (('summary','Краткое содержание'),('notes','Восстановленный конспект'),('definitions','Определения и ключевые понятия'),('formulas','Формулы и обозначения'),('emphasis','Акценты преподавателя'),('uncertainties','Сомнительные фрагменты')):
  d.add_paragraph(label,'Heading 1')
  if not c[k]:d.add_paragraph('Отдельные элементы по доступной расшифровке не установлены.')
  for x in c[k]:block(x)
 d.add_paragraph('Словарный пакет','Heading 1')
 p=d.add_paragraph('Термины этой лекции находятся в соседнем файле ')
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),d.part.relate_to('vocabulary_batch.json',RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='vocabulary_batch.json';r.append(t);h.append(r);p._p.append(h)
 d.add_paragraph(f"В пакете {len(c['vocabulary'])} терминов. Количество встреч — буквальные вхождения в SRT, включая повторы ASR; оно не оценивает число произнесений в аудио. Пакет не применён к общему словарю автоматически.")
 d.add_paragraph('Источники и границы восстановления','Heading 1')
 d.add_paragraph('Использованы только объединённый raw, все части SRT и manifest. Слайды, фотографии и аудио в эту обработку не включены. Учебная реконструкция не заменяет проверку повреждённых мест по записи.')
 for r in c['part_review']:d.add_paragraph(f"Часть {r['part']}. {r['note']}")
 footer=s.footer.paragraphs[0];footer.add_run('Lecture Pipeline 1.5 · ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 d.core_properties.title=c['title'];d.core_properties.subject=m['subject'];d.core_properties.author='Lecture Pipeline';d.save(dest)
def verify_inputs(b):
 for p,h in b['input_hashes'].items():need(Path(p).is_file() and sha(Path(p).read_bytes())==h,'Input changed: '+p)
def build(packet,content,out):
 started=time.perf_counter();b=load(packet);verify_inputs(b)
 # Do not trust packet's embedded segments: recompute them from hashed inputs.
 actual=read_bundle(b['manifest_path']);need(actual['input_hashes']==b['input_hashes'] and actual['segments']==b['segments'] and actual['manifest']==b['manifest'],'Tampered source packet')
 for key,value in {'subject_code':actual['manifest']['subject_code'],'subject':actual['manifest']['subject'],'lecture_id':actual['manifest']['lecture'],'source_file_name':Path(actual['raw_path']).name,'source_size_bytes':Path(actual['raw_path']).stat().st_size}.items():need(b['source'].get(key)==value,'Source metadata mismatch: '+key)
 provenance.verify(b['source'],actual)
 data,s=state_read(b['ai_state_path']);need(sha(data)==b['ai_state_sha256'],'ai_state changed since prepare')
 need(b['pilot'] or not already(s,provenance.key(b['source']),b['manifest']),'Already processed lecture')
 c=validate_content(load(content),b);dest=Path(out).resolve();need(not dest.exists(),'Refusing to replace outputs')
 dest.parent.mkdir(parents=True,exist_ok=True);stage=dest.parent/('.ai-v15-'+uuid.uuid4().hex);stage.mkdir()
 try:
  render_word(c,b,stage/'lecture.docx');terms=[]
  for t in c['vocabulary']:
   term={k:v for k,v in t.items() if k in ('chinese','pinyin','russian','note','entry_id','sense_id')}
   term['occurrences']=sum(s['text'].count(t['chinese']) for s in b['segments']);terms.append(term)
  batch=dict(schema_version=1,source=b['source'],terms=terms);dump(stage/'vocabulary_batch.json',batch)
  entry=dict(b['source'],lecture_date=b['manifest']['lecture'].split('_',1)[1],recording_parts=len(b['manifest']['outputs']),output_docx_name='lecture.docx',vocabulary_entries=len(terms),processed_at=utc(),word_ai_version=VERSION,vocabulary_status='batch_pending',output_docx_local_path=str(dest/'lecture.docx'))
  patch=dict(schema_version=1,expected_ai_state_sha256=b['ai_state_sha256'],entry=entry,requires_real_cloud_receipt=b['source'].get('provenance',{}).get('kind')!='local',pilot=b['pilot']);patch.update({'source_key':provenance.key(b['source'])} if 'source_key' in b['source'] else {'source_file_id':provenance.key(b['source'])});dump(stage/'ai_state_patch.json',patch)
  shutil.copy2(content,stage/'ai_content.json');shutil.copy2(packet,stage/'source_packet.json')
  hashes={p.name:sha(p.read_bytes()) for p in stage.iterdir() if p.is_file()}
  report=dict(schema_version=1,layer_version=VERSION,status='pilot_review_required' if b['pilot'] else 'review_required',pilot=b['pilot'],source=b['source'],input_hashes=b['input_hashes'],output_hashes=hashes,sections={k:len(c[k]) for k in ('summary','notes','definitions','formulas','emphasis','uncertainties')},organization={k:c['organization'][k]['status'] for k in ORG},part_review=c['part_review'],vocabulary_terms=len(terms),elapsed_build_seconds=round(time.perf_counter()-started,3),ai_generation='ChatGPT assisted structured content; no autonomous API call',validation='structure, hashes, literal quotes and SRT references; semantic correctness requires review',audio_listened=False,state_written=False,dictionary_written=False)
  dump(stage/'processing_report.json',report);verify_inputs(b);need(sha(Path(b['ai_state_path']).read_bytes())==b['ai_state_sha256'],'State changed during build');stage.rename(dest)
 except BaseException:
  shutil.rmtree(stage);raise
 return report
@contextmanager
def locked_file(path):
 with atomic_publish.exclusive(path) as f:yield f

def validate_state(data):
 s=json.loads(data.decode('utf-8-sig'))
 need(s.get('schema_version')==1 and isinstance(s.get('processed'),dict),'Unsupported ai_state schema')

def validate_docx(path):
 from docx import Document
 d=Document(path);texts={p.text for p in d.paragraphs}
 for heading in ('Организационные вопросы','Краткое содержание','Восстановленный конспект','Определения и ключевые понятия','Формулы и обозначения','Акценты преподавателя','Сомнительные фрагменты','Словарный пакет'):
  need(heading in texts,'Missing Word section: '+heading)
 return True

def commit(output,state,receipt,backup_dir):
 out=Path(output);report=load(out/'processing_report.json');patch=load(out/'ai_state_patch.json');review=load(receipt)
 need(not report['pilot'] and not patch['pilot'],'Pilot can never update ai_state')
 need(review.get('report_sha256')==sha((out/'processing_report.json').read_bytes()),'Review does not match report')
 local_mode=patch['entry'].get('provenance',{}).get('kind')=='local'
 external=provenance.drive_fields(review,required=not local_mode)
 if local_mode:
  need(review.get('local_validation')=='schema_refs_docx_v1','Local structural validation receipt required')
  validate_docx(out/'lecture.docx')
  need(review.get('source_key')==provenance.key(patch['entry']),'Receipt local source mismatch')
 else:
  need(review.get('content_reviewed') is True and review.get('all_pages_visually_checked') is True,'Content and visual review required')
  need(review['source_file_id']==patch['source_file_id'],'Receipt source mismatch')
 for name,h in report['output_hashes'].items():need(sha(local_file(out,name).read_bytes())==h,'Output changed after review: '+name)
 b=load(out/'source_packet.json');verify_inputs(b);provenance.verify(b['source'],read_bundle(b['manifest_path']));need(all(patch['entry'].get(k)==v for k,v in b['source'].items()),'Patch source mismatch')
 need(Path(state).resolve()==Path(b['ai_state_path']).resolve(),'Wrong ai_state file')
 def build_state(old):
  s=json.loads(old.decode('utf-8-sig'))
  need(s.get('schema_version')==1 and not already(s,provenance.key(patch['entry']),b['manifest']),'Already processed/unsupported state')
  verify_inputs(b)
  for name,h in report['output_hashes'].items():need(sha(local_file(out,name).read_bytes())==h,'Output changed after review: '+name)
  entry=copy.deepcopy(patch['entry']);entry.update(external)
  if local_mode:entry['output_provenance']=dict(kind='local',path=str((out/'lecture.docx').resolve()),sha256=sha((out/'lecture.docx').read_bytes()),review='schema/refs/DOCX structural validation; not human visual review')
  new=copy.deepcopy(s);new['processed'][provenance.key(patch['entry'])]=entry;new['updated_at']=utc()
  return (json.dumps(new,ensure_ascii=False,indent=2,allow_nan=False)+'\n').encode('utf8')
 transaction=atomic_publish.publish(state,patch['expected_ai_state_sha256'],build_state,validate_state,backup_dir)
 return dict(status='committed',backup=str(Path(transaction['backup']).parent),dictionary_status='batch_pending')

def main():
 p=argparse.ArgumentParser(description=__doc__);sub=p.add_subparsers(dest='command',required=True)
 a=sub.add_parser('prepare');a.add_argument('--manifest',required=True);a.add_argument('--state',required=True);a.add_argument('--source-id');a.add_argument('--workspace');a.add_argument('--out',required=True);a.add_argument('--pilot',action='store_true')
 a=sub.add_parser('build');a.add_argument('--packet',required=True);a.add_argument('--content',required=True);a.add_argument('--out',required=True)
 a=sub.add_parser('commit');a.add_argument('--output',required=True);a.add_argument('--state',required=True);a.add_argument('--receipt',required=True);a.add_argument('--backup-dir',required=True)
 args=vars(p.parse_args());command=args.pop('command');args={k.replace('-','_'):v for k,v in args.items()}
 result=globals()[command](**args);print(json.dumps(result if command!='prepare' else {'status':'prepared','parts':len(result['manifest']['outputs'])},ensure_ascii=False,indent=2))
if __name__=='__main__':main()
